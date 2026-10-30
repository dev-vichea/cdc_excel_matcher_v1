import unittest
from pathlib import Path
import openpyxl
from io import BytesIO

from services.cutstock_updater import (
    detect_period_columns,
    suggest_period_from_date,
    consolidate_deductions,
    apply_deductions_to_cdc_workbook,
)

ROOT = Path(__file__).resolve().parent.parent
EXAMPLE_2025 = ROOT / "examples" / "cdc_master_2025.xlsx"
EXAMPLE_2026 = (
    ROOT
    / "reference"
    / "shipping_work"
    / "03_cdc_master_and_approval"
    / "needs_verification"
    / "XIN CAI FU-CDC Master List-2026-CUT STOCK ( QIP # 211-26  ).xlsx"
)


class TestCutStockUpdater(unittest.TestCase):

    def test_detect_periods_english_2026(self):
        if not EXAMPLE_2026.exists():
            self.skipTest("Reference file 2026 not found")
        data = EXAMPLE_2026.read_bytes()
        periods = detect_period_columns(data)
        self.assertTrue(len(periods) >= 12, f"Expected >= 12 periods, got {len(periods)}")
        labels = [p["label"] for p in periods]
        self.assertIn("AUG 2026", labels)
        self.assertIn("SEP 2026", labels)
        self.assertIn("OCT 2026", labels)

        # Check column index for SEP 2026 is Col L (12)
        sep = next(p for p in periods if p["label"] == "SEP 2026")
        self.assertEqual(sep["col_idx"], 12)
        self.assertEqual(sep["col_letter"], "L")

    def test_detect_periods_khmer_2025(self):
        if not EXAMPLE_2025.exists():
            self.skipTest("Example 2025 file not found")
        data = EXAMPLE_2025.read_bytes()
        periods = detect_period_columns(data)
        self.assertTrue(len(periods) >= 4, f"Expected >= 4 periods, got {len(periods)}")
        labels = [p["label"] for p in periods]
        self.assertTrue(any("ត្រីមាសទី១" in lbl for lbl in labels))
        # Check first quarter is Col G (7)
        q1 = periods[0]
        self.assertEqual(q1["col_idx"], 7)
        self.assertEqual(q1["col_letter"], "G")
        self.assertEqual(q1["amount_col_idx"], 8)

    def test_suggest_period_from_date(self):
        periods = [
            {"col_idx": 11, "label": "AUG 2026"},
            {"col_idx": 12, "label": "SEP 2026"},
            {"col_idx": 14, "label": "OCT 2026"},
        ]
        col = suggest_period_from_date("2026.9.20", periods)
        self.assertEqual(col, 12)
        col_oct = suggest_period_from_date("2026-10-08", periods)
        self.assertEqual(col_oct, 14)

    def test_consolidate_deductions_multiline(self):
        results = [
            {"row": 1, "chinese": "堆垛架", "english_provided": "Fabric roll rack", "qty": 15.0, "unit": "PCS", "unit_price": 128.0, "total": 1920.0},
            {"row": 2, "chinese": "周转车", "english_provided": "Fabric roll rack", "qty": 14.0, "unit": "PCS", "unit_price": 147.0, "total": 2058.0},
            {"row": 3, "chinese": "周转车", "english_provided": "Fabric roll rack", "qty": 27.0, "unit": "PCS", "unit_price": 110.0, "total": 2970.0},
            {"row": 4, "chinese": "工业缝纫机", "english_provided": "Auto template machine", "qty": 6.0, "unit": "SET", "unit_price": 4800.0, "total": 28800.0},
        ]
        selections = [
            {"match": {"sheet": "Sheet1", "row": 78, "serial": 71, "description": "Fabric roll rack", "unit": "Set", "approved": 5000, "remaining": 5000, "unit_compatible": True}},
            {"match": {"sheet": "Sheet1", "row": 78, "serial": 71, "description": "Fabric roll rack", "unit": "Set", "approved": 5000, "remaining": 5000, "unit_compatible": True}},
            {"match": {"sheet": "Sheet1", "row": 78, "serial": 71, "description": "Fabric roll rack", "unit": "Set", "approved": 5000, "remaining": 5000, "unit_compatible": True}},
            {"match": {"sheet": "Sheet1", "row": 96, "serial": 89, "description": "Auto template machine", "unit": "Set", "approved": 10, "remaining": 10, "unit_compatible": True}},
        ]
        res = consolidate_deductions(results, selections)
        self.assertEqual(len(res["consolidated"]), 2)
        self.assertFalse(res["has_exceeded"])

        rack = next(c for c in res["consolidated"] if c["cdc_row"] == 78)
        self.assertEqual(rack["deduct_qty"], 56.0)
        self.assertEqual(rack["total_invoice_usd"], 6948.0)
        self.assertEqual(rack["new_remaining"], 4944.0)
        self.assertFalse(rack["is_exceeded"])

        auto = next(c for c in res["consolidated"] if c["cdc_row"] == 96)
        self.assertEqual(auto["deduct_qty"], 6.0)
        self.assertEqual(auto["new_remaining"], 4.0)

    def test_apply_deductions_preserves_layout_and_formulas_2026(self):
        if not EXAMPLE_2026.exists():
            self.skipTest("Reference file 2026 not found")
        orig_data = EXAMPLE_2026.read_bytes()
        deductions = [
            {"cdc_sheet": "Sheet1", "cdc_row": 78, "deduct_qty": 56.0, "total_invoice_usd": 6948.0},
            {"cdc_sheet": "Sheet1", "cdc_row": 96, "deduct_qty": 6.0, "total_invoice_usd": 28800.0},
        ]
        # Col 12 is SEP 2026
        updated_bytes = apply_deductions_to_cdc_workbook(orig_data, deductions, target_col_idx=12)

        # Inspect updated workbook
        wb = openpyxl.load_workbook(BytesIO(updated_bytes), data_only=False)
        ws = wb.active

        # Check headers intact
        self.assertEqual(ws.cell(1, 1).value, "CIB 01E(i)")
        self.assertIn("NEW FOCUS", str(ws.cell(1, 6).value))
        self.assertEqual(ws.cell(5, 1).value, "ITEM NO")
        self.assertEqual(ws.cell(5, 2).value, "DESCRIPTION OF GOODS")

        # Check values written
        self.assertEqual(ws.cell(78, 12).value, 56.0)
        self.assertEqual(ws.cell(96, 12).value, 6.0)

        # Check formulas preserved
        formula_78 = ws.cell(78, 29).value  # Col AC ending balance
        self.assertTrue(isinstance(formula_78, str) and formula_78.startswith("=G78-"))
        amount_formula_78 = ws.cell(78, 30).value  # Col AD balance amount
        self.assertEqual(amount_formula_78, "=AC78*H78")

    def test_apply_deductions_khmer_2025(self):
        if not EXAMPLE_2025.exists():
            self.skipTest("Example 2025 file not found")
        orig_data = EXAMPLE_2025.read_bytes()
        deductions = [
            {"cdc_sheet": "Sheet1", "cdc_row": 13, "deduct_qty": 10.0, "total_invoice_usd": 4000.0, "amount_col_idx": 8},
        ]
        updated_bytes = apply_deductions_to_cdc_workbook(orig_data, deductions, target_col_idx=7)

        wb = openpyxl.load_workbook(BytesIO(updated_bytes), data_only=False)
        ws = wb.active

        # Check Khmer header intact
        self.assertEqual(ws.cell(1, 14).value, "ព្រះរាជាណាចក្រកម្ពុជា")
        self.assertIn("KING SG APPAREL", str(ws.cell(2, 1).value))

        # Check quantity written
        self.assertEqual(ws.cell(13, 7).value, 10.0)
        # Check formula in Col 8 (Amount) was PRESERVED
        self.assertEqual(ws.cell(13, 8).value, "=G13*E13")
        # Check formula in Col 17 (Remaining Qty) was PRESERVED
        self.assertEqual(ws.cell(13, 17).value, "=D13-G13-I13-K13-M13")


if __name__ == "__main__":
    unittest.main()
