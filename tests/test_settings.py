"""Tests for Settings & Preferences store."""

import unittest
from pathlib import Path
from services.settings_store import (
    get_settings, save_settings,
    get_default_template_path, get_default_cdc_path,
)
from services.extraction import extract_cdc


class TestSettingsStore(unittest.TestCase):
    def test_settings_retrieval(self):
        settings = get_settings()
        self.assertIn("company_name", settings)
        self.assertIn("qip_reference", settings)
        self.assertIn("vat_number", settings)
        self.assertIn("shipping_mode", settings)
        self.assertEqual(settings["qip_reference"], "211/26")

    def test_default_template_path_exists(self):
        p = get_default_template_path()
        self.assertTrue(p.exists(), f"Default template path {p} does not exist")
        self.assertTrue(p.name.endswith(".xlsx"))

    def test_default_cdc_path_exists_and_extracts(self):
        p = get_default_cdc_path()
        self.assertTrue(p.exists(), f"Default CDC path {p} does not exist")
        meta, items = extract_cdc(p.read_bytes(), p.name)
        self.assertGreater(len(items), 50)
        # Check that unit price is extracted for items
        items_with_price = [it for it in items if it.get("unit_price") is not None]
        self.assertGreater(len(items_with_price), 0)

    def test_auto_load_defaults_end_to_end(self):
        """Verify that default CDC master and default template auto-load and process without manual uploads."""
        from services.settings_store import get_default_cdc_bytes, get_default_template_bytes
        from services.extraction import extract_shipping, extract_cdc
        from services.packing import extract_packing, reconcile_invoice_packing
        from services.matching import build_results, load_mappings
        from services.permit_exporter import generate_permit_document
        import io
        import openpyxl

        # 1. User provides only supplier invoice
        ship_path = Path(__file__).parent.parent / 'reference' / 'shipping_work' / '02_import_TK164' / 'supplier_originals' / 'RE2612040-SHIPPING DOCS-28-OCT.xls'
        ship_bytes = ship_path.read_bytes()

        # 2. CDC master & template are auto-loaded from local machine
        cdc_bytes = get_default_cdc_bytes()
        tpl_bytes = get_default_template_bytes()

        smeta, sitems = extract_shipping(ship_bytes, ship_path.name)
        pmeta, pitems = extract_packing(ship_bytes, ship_path.name)
        cmeta, citems = extract_cdc(cdc_bytes, "default_cdc_master.xlsx")
        pairs, issues = reconcile_invoice_packing(sitems, pitems, smeta, pmeta)

        mappings = load_mappings(Path(__file__).parent.parent / "data" / "approved_mappings.json")
        results = build_results(sitems, citems, mappings, online=False)
        self.assertEqual(len(results), 10)

        selections = [{'match': r['candidates'][0] if r['candidates'] else None, 'decision': 'Approved', 'notes': ''} for r in results]

        # 3. Generate permit document using auto-loaded template
        permit_bytes = generate_permit_document(
            smeta, cmeta, results, selections, pitems, pairs, issues, '211/26',
            template_bytes=tpl_bytes
        )
        wb = openpyxl.load_workbook(io.BytesIO(permit_bytes))
        self.assertIn("INV", wb.sheetnames)
        self.assertIn("PL", wb.sheetnames)
        # Verify INV items start at row 24
        self.assertEqual(wb['INV'].cell(24, 1).value, 4)
        self.assertEqual(wb['INV'].cell(24, 6).value, 4.5)  # CDC unit price


if __name__ == "__main__":
    unittest.main()
