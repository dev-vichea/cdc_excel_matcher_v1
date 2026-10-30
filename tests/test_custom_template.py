import io
import unittest
from pathlib import Path
import openpyxl

from services.extraction import extract_shipping, extract_cdc
from services.packing import extract_packing, reconcile_invoice_packing
from services.matching import build_results, load_mappings
from services.permit_exporter import generate_permit_document

ROOT = Path(__file__).parent.parent


class TestCustomTemplate(unittest.TestCase):
    def setUp(self):
        self.re26_path = ROOT / 'reference' / 'shipping_work' / '02_import_TK164' / 'supplier_originals' / 'RE2612040-SHIPPING DOCS-28-OCT.xls'
        self.cdc_path = ROOT / 'reference' / 'shipping_work' / '03_cdc_master_and_approval' / 'needs_verification' / 'XIN CAI FU-CDC Master List-2026-CUT STOCK ( QIP # 211-26  ).xlsx'
        self.default_tpl_path = ROOT / 'reference' / 'shipping_work' / '02_import_TK164' / 'permit_application' / '1-26TKYM164.xlsx'

        if not self.re26_path.exists() or not self.cdc_path.exists():
            self.skipTest('Reference files missing')

        ship_bytes = self.re26_path.read_bytes()
        cdc_bytes = self.cdc_path.read_bytes()

        self.smeta, self.sitems = extract_shipping(ship_bytes, self.re26_path.name)
        self.pmeta, self.pitems = extract_packing(ship_bytes, self.re26_path.name)
        self.cmeta, self.citems = extract_cdc(cdc_bytes, self.cdc_path.name)
        self.pairs, self.issues = reconcile_invoice_packing(self.sitems, self.pitems, self.smeta, self.pmeta)

        mappings = load_mappings(ROOT / 'data' / 'approved_mappings.json')
        self.results = build_results(self.sitems, self.citems, mappings, online=False)
        self.selections = [{'match': r['candidates'][0] if r['candidates'] else None, 'decision': 'Approved', 'notes': ''} for r in self.results]

    def test_custom_template_renamed_sheets(self):
        """Test that uploaded template with custom sheet names ('Invoice', 'Packing List') works seamlessly."""
        wb = openpyxl.load_workbook(self.default_tpl_path)
        wb['INV'].title = 'Invoice'
        wb['PL'].title = 'Packing List'

        # Also put a custom marker in cell A5 to ensure uploaded template is the one used
        wb['Invoice']['A5'] = 'CUSTOM TEMPLATE TEST MARKER'

        bio = io.BytesIO()
        wb.save(bio)
        custom_bytes = bio.getvalue()

        out_bytes = generate_permit_document(
            self.smeta, self.cmeta, self.results, self.selections,
            self.pitems, self.pairs, self.issues, '211/26',
            template_bytes=custom_bytes
        )

        out_wb = openpyxl.load_workbook(io.BytesIO(out_bytes))
        self.assertIn('Invoice', out_wb.sheetnames)
        self.assertIn('Packing List', out_wb.sheetnames)

        inv = out_wb['Invoice']
        pl = out_wb['Packing List']

        # Verify custom template marker was preserved
        self.assertEqual(inv['A5'].value, 'CUSTOM TEMPLATE TEST MARKER')

        # Verify populated fields
        self.assertEqual(inv['B8'].value, 'RE2612040')
        self.assertEqual(pl['B8'].value, 'RE2612040')  # Resolved static string
        self.assertEqual(inv.cell(24, 1).value, 4)     # CDC #4 (row 24 in new template)
        self.assertEqual(inv.cell(29, 4).value, 407)   # Elastic band converted qty (row 29)

    def test_single_sheet_custom_template(self):
        """Test that a template with only an Invoice sheet works without raising errors."""
        wb = openpyxl.load_workbook(self.default_tpl_path)
        del wb['PL']  # Delete PL sheet

        bio = io.BytesIO()
        wb.save(bio)
        single_sheet_bytes = bio.getvalue()

        out_bytes = generate_permit_document(
            self.smeta, self.cmeta, self.results, self.selections,
            self.pitems, self.pairs, self.issues, '211/26',
            template_bytes=single_sheet_bytes
        )

        out_wb = openpyxl.load_workbook(io.BytesIO(out_bytes))
        self.assertIn('INV', out_wb.sheetnames)
        self.assertEqual(len(out_wb.sheetnames), 1)
        self.assertEqual(out_wb['INV']['B8'].value, 'RE2612040')


if __name__ == '__main__':
    unittest.main()
