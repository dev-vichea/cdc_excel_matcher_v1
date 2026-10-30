import unittest
from pathlib import Path
import importlib.util
import tempfile
import json
import io
from unittest.mock import patch
from openpyxl import load_workbook
from services.extraction import extract_shipping,extract_cdc
from services.packing import extract_packing,reconcile_invoice_packing
from services.matching import build_results,translate
from services.translation_memory import load_verified_translations,save_verified_translation,migrate_json_translations
from services.exporter import export_report
from services.permit_exporter import export_permit_draft

ROOT=Path(__file__).parents[1]
class MatchingTests(unittest.TestCase):
    def test_dictionary(self):
        self.assertEqual(translate('堆垛架')[0],'stacking rack')
    def test_verified_shipping_term_reused_for_matching(self):
        ship={'sheet':'invoice','row':14,'chinese':'新布架','english_provided':'Storage shelf',
              'qty':1,'unit':'Set'}
        cdc=[{'sheet':'CDC','row':8,'description':'Storage shelf','unit':'Set','approved':1,'remaining':1},
             {'sheet':'CDC','row':9,'description':'Fabric rack','unit':'Set','approved':1,'remaining':1}]
        before=build_results([ship],cdc)[0]
        self.assertEqual(before['candidates'][0]['description'],'Storage shelf')
        with tempfile.TemporaryDirectory() as directory:
            path=Path(directory)/'translations.sqlite3'
            save_verified_translation(path,'新布架','Fabric rack')
            glossary=load_verified_translations(path)
            after=build_results([ship],cdc,glossary=glossary)[0]
        self.assertEqual(after['translated'],'Fabric rack')
        self.assertEqual(after['translation_method'],'saved translation')
        self.assertEqual(after['candidates'][0]['description'],'Fabric rack')
        corrected=build_results([ship],cdc,overrides={('invoice',14):'Fabric rack'})[0]
        self.assertEqual(corrected['translation_method'],'reviewer correction')
        self.assertEqual(corrected['candidates'][0]['description'],'Fabric rack')
    def test_json_translation_migration_preserves_later_sqlite_edit(self):
        with tempfile.TemporaryDirectory() as directory:
            json_path=Path(directory)/'translations.json'
            database_path=Path(directory)/'translations.sqlite3'
            json_path.write_text(json.dumps({'新布架':'Fabric rack'},ensure_ascii=False),encoding='utf-8')
            self.assertEqual(migrate_json_translations(json_path,database_path),1)
            self.assertEqual(load_verified_translations(database_path)['新布架'],'Fabric rack')
            save_verified_translation(database_path,'新布架','Roll rack')
            self.assertEqual(migrate_json_translations(json_path,database_path),0)
            self.assertEqual(load_verified_translations(database_path)['新布架'],'Roll rack')
    @unittest.skipUnless(importlib.util.find_spec('xlrd'), 'xlrd is installed via requirements.txt for actual XLS')
    def test_actual_documents(self):
        ship=(ROOT/'examples/shipping_TK164.xls')
        cdc=(ROOT/'examples/cdc_master_2025.xlsx')
        meta,items=extract_shipping(ship.read_bytes(),ship.name)
        info,cdc_items=extract_cdc(cdc.read_bytes(),cdc.name)
        self.assertEqual(meta['invoice_no'],'26TKYM164')
        self.assertEqual(len(items),4)
        self.assertEqual([i['qty'] for i in items],[15.,14.,27.,6.])
        self.assertTrue(len(cdc_items)>20)
        self.assertIn('KING SG',info['company'].upper())
        candidates=build_results(items,cdc_items)
        self.assertEqual(len(candidates),4)
        output=export_report(meta,info,candidates,[{'match':None,'decision':'Pending review','notes':''} for _ in candidates])
        self.assertTrue(output.startswith(b'PK'))
    def test_supplier_packing_reconciliation_and_permit_draft(self):
        source=ROOT/'examples/shipping_TK164.xls'
        data=source.read_bytes()
        invoice_meta,invoice_items=extract_shipping(data,source.name)
        packing_meta,packing_items=extract_packing(data,source.name)
        pairs,issues=reconcile_invoice_packing(invoice_items,packing_items,invoice_meta,packing_meta)
        self.assertEqual(len(packing_items),4)
        self.assertEqual(issues,[])
        self.assertEqual([p['packing']['row'] for p in pairs],[13,14,15,16])
        self.assertEqual(packing_meta['printed_totals']['packages'],73.)
        self.assertEqual(packing_meta['printed_totals']['gross_weight'],7220.)
        self.assertEqual(packing_meta['printed_totals']['volume'],43.84)
        altered=[dict(item) for item in packing_items]
        altered[1]['qty']=13.
        _,changed=reconcile_invoice_packing(invoice_items,altered,invoice_meta,packing_meta)
        self.assertTrue(any('quantity differs' in issue for issue in changed))
        altered=[dict(item) for item in packing_items]
        altered[1]['gross_weight']=None
        _,changed=reconcile_invoice_packing(invoice_items,altered,invoice_meta,packing_meta)
        self.assertTrue(any('missing gross weight' in issue for issue in changed))
        results=build_results(invoice_items,[])
        selections=[{'match':None,'decision':'Pending review','notes':''} for _ in results]
        selected={'serial':'96','description':'Auto template machine','unit':'SET',
                  'approved':10.,'remaining':None,'sheet':'CDC master','row':96,'score':100.,
                  'unit_compatible':True}
        selections[-1]={'match':selected,'decision':'Reviewed candidate (not approved)','notes':''}
        output=export_permit_draft(invoice_meta,{'company':invoice_meta['consignee'],'qip':''},
                                   results,selections,packing_items,pairs,issues,'211/26')
        workbook=load_workbook(io.BytesIO(output),read_only=True,data_only=True)
        self.assertEqual(workbook.sheetnames,['INV','PL'])
        self.assertEqual(workbook['INV']['A15'].value,'#96')
        self.assertEqual(workbook['INV']['D16'].value,62.)
        self.assertEqual(workbook['INV']['G16'].value,35748.)
        self.assertEqual(workbook['PL']['A16'].value,73.)
        self.assertEqual(workbook['PL']['F16'].value,7220.)
        self.assertEqual(workbook['PL']['G16'].value,43.84)
        self.assertNotIn('KING SG',str(workbook['INV']['B4'].value))
        workbook.close()
        missing_amount=[dict(result) for result in results]
        missing_amount[0]['total']=None
        with self.assertRaises(ValueError):
            export_permit_draft(invoice_meta,{'company':invoice_meta['consignee']},
                                missing_amount,selections,packing_items,pairs,issues,'211/26')
    def test_2025_cdc_and_export(self):
        cdc=ROOT/'examples/cdc_master_2025.xlsx'
        info,items=extract_cdc(cdc.read_bytes(),cdc.name)
        self.assertGreater(len(items),50)
        self.assertIn('KING SG',info['company'].upper())
        r={'sheet':'invoice','row':17,'hs_code':'8452290000','chinese':'工业缝纫机','english_provided':'Auto template machine','qty':6.0,'unit':'SET','unit_price':4800.,'total':28800.}
        results=build_results([r],items)
        self.assertTrue(results[0]['candidates'])
        blob=export_report({'invoice_no':'26TKYM164'},info,results,[{'match':None,'decision':'Pending review','notes':''}])
        self.assertTrue(blob.startswith(b'PK'))
    def test_2026_cdc_layout(self):
        header=[None]*29
        header[1]='DESCRIPTION OF GOODS'; header[5]='UNIT'; header[6]='QUANTITY'; header[28]='ENDING'
        section=['I','MACHINERIES AND EQUIPMENT']+[None]*27
        item=[1,'One needle sewing machine',None,None,'New','Set',400]+[None]*21+[386]
        rows=[(1,[None]*5+['COMPANY   : NEW FOCUS APPAREL (CAMBODIA) Co., LTD.']),
              (5,header),(7,section),(8,item)]
        with patch('services.extraction.rows_from_workbook',return_value=iter([('Sheet1',rows)])):
            info,items=extract_cdc(b'', '2026.xlsx')
        self.assertEqual(info['company'],'NEW FOCUS APPAREL (CAMBODIA) Co., LTD.')
        self.assertEqual(len(items),1)
        self.assertEqual((items[0]['description'],items[0]['unit'],items[0]['approved'],items[0]['remaining']),
                         ('One needle sewing machine','Set',400.,386.))
    def test_2026_application_master_without_balance(self):
        header=[None]*7
        header[1]='Descriptions of Goods'; header[5]='Unit'; header[6]='Quantity'
        item=[1,'One needle sewing machine',None,None,'New','Set',400]
        rows=[(1,[None]*3+['Company : XIN CAI FU (CAMBODIA) CO., LTD. (ស៊ីន កៃ ហូ)']),
              (5,header),(7,['I','MACHINERIES AND EQUIPMENT']),(8,item)]
        with patch('services.extraction.rows_from_workbook',return_value=iter([('CDC master',rows)])):
            info,items=extract_cdc(b'', 'master.xlsx')
        self.assertEqual(info['company'],'XIN CAI FU (CAMBODIA) CO., LTD.')
        self.assertEqual(len(items),1)
        self.assertEqual(items[0]['approved'],400.)
        self.assertIsNone(items[0]['remaining'])
if __name__=='__main__':unittest.main()
