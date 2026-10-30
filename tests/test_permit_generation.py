import io
import unittest
from pathlib import Path
import openpyxl

from services.extraction import extract_shipping, extract_cdc
from services.packing import extract_packing, reconcile_invoice_packing
from services.matching import build_results, load_mappings
from services.translation_memory import load_verified_translations
from services.permit_exporter import generate_permit_document
from services.cutstock_updater import (
    detect_period_columns,
    suggest_period_from_date,
    consolidate_deductions,
    apply_deductions_to_cdc_workbook,
)

ROOT = Path(__file__).parent.parent


class TestPermitGeneration(unittest.TestCase):
    def setUp(self):
        self.re26_path = ROOT / 'reference' / 'shipping_work' / '02_import_TK164' / 'supplier_originals' / 'RE2612040-SHIPPING DOCS-28-OCT.xls'
        self.cdc_path = ROOT / 'reference' / 'shipping_work' / '03_cdc_master_and_approval' / 'needs_verification' / 'XIN CAI FU-CDC Master List-2026-CUT STOCK ( QIP # 211-26  ).xlsx'

    def test_re2612040_full_pipeline(self):
        if not self.re26_path.exists() or not self.cdc_path.exists():
            self.skipTest('Reference files not found')

        ship_bytes = self.re26_path.read_bytes()
        cdc_bytes = self.cdc_path.read_bytes()

        # 1. Extraction
        smeta, sitems = extract_shipping(ship_bytes, self.re26_path.name)
        pmeta, pitems = extract_packing(ship_bytes, self.re26_path.name)
        cmeta, citems = extract_cdc(cdc_bytes, self.cdc_path.name)

        self.assertEqual(len(sitems), 10)
        self.assertEqual(len(pitems), 10)
        self.assertEqual(smeta['invoice_no'], 'RE2612040')
        self.assertEqual(smeta['invoice_date'], '2026-10-28')
        self.assertIn('WEALDONE', smeta['supplier'])
        self.assertIn('XIN CAI FU', smeta['consignee'])

        # 2. Packing reconciliation
        pairs, issues = reconcile_invoice_packing(sitems, pitems, smeta, pmeta)
        self.assertEqual(issues, [])
        self.assertEqual(sum(p['packing']['packages'] for p in pairs), 390.0)

        # 3. Matching
        mappings = load_mappings(ROOT / 'data' / 'approved_mappings.json')
        trans = load_verified_translations(ROOT / 'data' / 'verified_translations.sqlite3')
        results = build_results(sitems, citems, mappings, online=False, glossary=trans)
        self.assertEqual(len(results), 10)

        selections = []
        for r in results:
            cand = r['candidates'][0] if r['candidates'] else None
            selections.append({'match': cand, 'decision': 'Approved', 'notes': ''})

        # Verify all 10 matched
        matched = [s['match'] for s in selections if s['match']]
        self.assertEqual(len(matched), 10)

        # 4. Generate Permit Application Document (1-26TKYM164 format)
        # 4a. With Bill of Lading metadata
        from services.bill_extractor import extract_bill
        bill_path = ROOT / 'reference' / 'shipping_work' / '02_import_TK164' / 'transport' / '汇孚390件提单AOTAI FAUXFUR.pdf'
        bill_meta = extract_bill(bill_path.read_bytes(), bill_path.name) if bill_path.exists() else {}

        permit_bytes = generate_permit_document(
            smeta, cmeta, results, selections, pitems, pairs, issues, '211/26',
            bill_meta=bill_meta
        )
        wb_permit = openpyxl.load_workbook(io.BytesIO(permit_bytes))
        self.assertIn('INV', wb_permit.sheetnames)
        self.assertIn('PL', wb_permit.sheetnames)

        inv = wb_permit['INV']
        pl = wb_permit['PL']

        self.assertEqual(inv['A1'].value, 'WEALDONE LIMITED')
        self.assertEqual(inv['B8'].value, 'RE2612040')
        # Bill date takes priority over current date!
        self.assertEqual(inv['G8'].value, '29-Oct-2026')
        self.assertEqual(inv['B9'].value, 'BY TRUCK')
        self.assertEqual(inv['B10'].value, 'JIANGSU, CHINA TO PHNOM PENH, CAMBODIA')
        self.assertEqual(pl['B10'].value, 'JIANGSU, CHINA TO PHNOM PENH, CAMBODIA')

        # Check Consignee address split across rows 13 and 14
        self.assertEqual(inv.cell(12, 2).value, 'XIN CAI FU (CAMBODIA) CO., LTD.')
        self.assertEqual(inv.cell(13, 2).value, 'PHUM DOMNAK TROYEOUNG, SANGKAT CHOM CHAO II,')
        self.assertEqual(inv.cell(14, 2).value, 'KHAN POR SEN CHEY PHNOM PENH,')
        self.assertEqual(inv.cell(15, 2).value, 'KINGDOM OF CAMBODIA.')
        self.assertEqual(inv.cell(16, 2).value, 'SHIPPING PERSON:')
        self.assertEqual(inv.cell(17, 2).value, 'VAT No : LO01-902303477')

        self.assertEqual(pl.cell(12, 2).value, 'XIN CAI FU (CAMBODIA) CO., LTD.')
        self.assertEqual(pl.cell(13, 2).value, 'PHUM DOMNAK TROYEOUNG, SANGKAT CHOM CHAO II,')
        self.assertEqual(pl.cell(14, 2).value, 'KHAN POR SEN CHEY PHNOM PENH,')
        self.assertEqual(pl.cell(15, 2).value, 'KINGDOM OF CAMBODIA.')
        self.assertEqual(pl.cell(16, 2).value, 'SHIPPING PERSON:')
        self.assertEqual(pl.cell(17, 2).value, 'VAT No : LO01-902303477')

        self.assertEqual(inv['A23'].value, '211/26')
        self.assertEqual(inv['B23'].value, 'II - RAW MATERIALS AND ACCESSORIES')

        # Check items and formulas (sorted smallest to largest by CDC No: 4, 13, 14, 17, 26, 28, 29, 31, 61, 86)
        # On INV: Items start at row 24
        self.assertEqual(inv.cell(24, 1).value, 4)  # Row 24 is CDC #4: Knitted fabric
        self.assertEqual(inv.cell(24, 2).value, 'KNITTED FABRIC')
        self.assertEqual(inv.cell(24, 4).value, 6406)
        self.assertEqual(inv.cell(24, 5).value, 'Kg')
        self.assertEqual(inv.cell(24, 5).alignment.horizontal, 'left')
        self.assertEqual(inv.cell(24, 6).value, 4.5)  # Matches CDC Master List unit price
        self.assertEqual(inv.cell(24, 7).value, '=D24*F24')

        self.assertEqual(pl['B8'].value, 'RE2612040')  # No raw =INV!B8 formula!
        self.assertEqual(pl['H8'].value, '29-Oct-2026')
        self.assertEqual(pl.cell(43, 2).value, '1X45 HQ PART OF CFS/CFS')

        # On PL: Items start at row 24 (row 23 is Category header)
        self.assertEqual(pl.cell(24, 1).value, 284)  # Row 24 is Knitted fabric (284 pkgs)
        self.assertEqual(pl.cell(24, 2).value, 'KNITTED FABRIC')

        self.assertEqual(inv.cell(25, 1).value, 13) # Row 25 is CDC #13: Ribbon
        self.assertEqual(inv.cell(25, 2).value, 'RIBBON')
        self.assertEqual(pl.cell(25, 1).value, 1)   # Ribbon has 1 pkg on PL (row 25)
        self.assertEqual(pl.cell(25, 2).value, 'RIBBON')
        self.assertEqual(pl.cell(25, 6).alignment.horizontal, 'left')

        # Row 29 on INV and PL is CDC #28: Elastic band (converted qty: 407 Gross, 36 pkgs)
        self.assertEqual(inv.cell(29, 1).value, 28)
        self.assertEqual(inv.cell(29, 2).value, 'ELASTIC BAND')
        self.assertEqual(inv.cell(29, 4).value, 407)
        self.assertEqual(inv.cell(29, 5).value, 'Gross')
        self.assertEqual(inv.cell(29, 6).value, 3.0)  # Matches CDC Master List unit price ($3.00, not invoice-derived $8.66)
        self.assertEqual(inv.cell(29, 7).value, '=D29*F29')

        self.assertEqual(pl.cell(29, 1).value, 36)
        self.assertEqual(pl.cell(29, 2).value, 'ELASTIC BAND')
        self.assertEqual(pl.cell(29, 5).value, 407)
        self.assertEqual(pl.cell(29, 6).value, 'Gross')

        self.assertEqual(inv.cell(33, 1).value, 86) # Row 33 is CDC #86: Elastic string
        self.assertEqual(inv.cell(40, 7).value, '=SUM(G24:G33)')
        self.assertEqual(inv.cell(43, 5).value, 'WEALDONE LIMITED')

        self.assertEqual(pl.cell(40, 1).value, '=SUM(A24:A33)')
        self.assertEqual(pl.cell(43, 6).value, 'WEALDONE LIMITED')

        # 5. Cut-Stock Deduction
        periods = detect_period_columns(cdc_bytes)
        sug_col = suggest_period_from_date(smeta['invoice_date'], periods)
        self.assertEqual(sug_col, 14)  # OCT 2026 is col 14 (N)

        consolidation = consolidate_deductions(results, selections)
        self.assertEqual(len(consolidation['consolidated']), 10)
        self.assertFalse(consolidation['has_exceeded'])

        # Verify Elastic band deduction is 407 Gross (not 58724 meters)
        elastic_cons = next(c for c in consolidation['consolidated'] if str(c['cdc_serial']) in ('28', 28))
        self.assertEqual(elastic_cons['deduct_qty'], 407.0)
        self.assertEqual(elastic_cons['cdc_unit'], 'Gross')

        deducted_bytes = apply_deductions_to_cdc_workbook(
            cdc_bytes,
            consolidation['consolidated'],
            target_col_idx=sug_col,
            sheet_name='Sheet1',
            add_to_existing=True
        )
        self.assertGreater(len(deducted_bytes), 1000)

        # Check deducted value in cell N104
        wb_deducted = openpyxl.load_workbook(io.BytesIO(deducted_bytes), data_only=True)
        ws_d = wb_deducted['Sheet1']
        self.assertEqual(ws_d.cell(104, 14).value, 6406.0)


if __name__ == '__main__':
    unittest.main()
