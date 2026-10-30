"""Produce new workbook only; never write to the source CDC workbook."""
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.utils import get_column_letter


def export_report(metadata,cdc_info,results,selections):
    wb=Workbook();ws=wb.active;ws.title='Matching Review'
    ws.append(['CDC SHIPMENT MATCHING — DRAFT ONLY / NO STOCK CHANGES'])
    ws.append(['Invoice',metadata.get('invoice_no',''),'Invoice consignee',metadata.get('consignee','')])
    ws.append(['CDC QIP',cdc_info.get('qip',''),'CDC company',cdc_info.get('company','')])
    ws.append(['NOTICE','Suggestions and user selections are NOT approved CDC deductions. Verify company/QIP, identity, units and eligibility.'])
    ws.append([])
    headers=['Invoice sheet','Invoice row','HS code','Chinese description','Provided English','Translation','Translation method',
             'Quantity','Ship unit','Unit price USD','Invoice value USD','Selected CDC sheet','Selected CDC row',
             'Selected CDC description','CDC unit','Approved qty','Available qty (if stored)','Similarity score',
             'Unit comparison','Review decision','Notes']
    ws.append(headers)
    opts=wb.create_sheet('Top Candidates')
    opts.append(['Invoice row','Chinese','Search description','Rank','CDC sheet','CDC row','CDC description','CDC unit','Similarity (0-100)','Unit comparison'])
    for n,r in enumerate(results):
        selection=selections[n];match=selection.get('match')
        decision=selection.get('decision','Pending review')
        ws.append([r['sheet'],r['row'],r['hs_code'],r['chinese'],r['english_provided'],r['translated'],r['translation_method'],
                   r['qty'],r['unit'],r['unit_price'],r['total'],match['sheet'] if match else '',match['row'] if match else '',
                   match['description'] if match else '',match['unit'] if match else '',match['approved'] if match else '',
                   match['remaining'] if match and match['remaining'] is not None else '',match['score'] if match else '',
                   ('Same text' if match['unit_compatible'] else 'CHECK UNIT') if match else '', decision,selection.get('notes','')])
        for rank,c in enumerate(r['candidates'],1):
            opts.append([r['row'],r['chinese'],r['query'],rank,c['sheet'],c['row'],c['description'],c['unit'],c['score'],
                         'Same text' if c['unit_compatible'] else 'CHECK UNIT'])
    summary=wb.create_sheet('Summary');summary.append(['Indicator','Value'])
    summary.append(['Invoice',metadata.get('invoice_no','')]);summary.append(['Shipping lines',len(results)])
    summary.append(['Selected candidates',sum(x.get('match') is not None for x in selections)])
    summary.append(['Needs review',sum(x.get('decision')!='Reviewed candidate (not approved)' for x in selections)])
    summary.append(['Stock deductions applied',0]);summary.append(['Original CDC workbook modified','NO'])
    for sh in wb:
        sh.freeze_panes='A7' if sh==ws else 'A2'
        header_row=6 if sh==ws else 1
        for cell in sh[header_row]:
            cell.fill=PatternFill('solid',fgColor='183153');cell.font=Font(color='FFFFFF',bold=True)
            cell.alignment=Alignment(wrap_text=True,vertical='center')
        sh.row_dimensions[header_row].height=34
        for col in sh.columns:
            letter=get_column_letter(col[0].column)
            width=max((len(str(c.value or '')) for c in col),default=8)
            sh.column_dimensions[letter].width=min(max(13,width+2),43)
        if sh==ws:
            sh.auto_filter.ref=f'A6:U{sh.max_row}'
            sh['A1'].font=Font(bold=True,color='9C2D19',size=14)
            sh['B4'].alignment=Alignment(wrap_text=True)
            sh.column_dimensions['D'].width=28
            sh.column_dimensions['N'].width=36
    out=BytesIO();wb.save(out);return out.getvalue()
