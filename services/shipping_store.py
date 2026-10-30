"""Local shipment register and document index for import and export work."""

from __future__ import annotations

from contextlib import contextmanager
from datetime import date, datetime
from hashlib import sha256
from pathlib import Path
import re
import sqlite3


IMPORT_DOCUMENTS = {
    "supplier_inv_pl": "Supplier invoice and packing list",
    "permit_inv_pl": "Reviewed permit invoice and packing list",
    "cdc_reference": "CDC master list / approval (if applicable)",
    "bill_of_lading": "Bill of lading",
    "arrival_notice": "Arrival notice / delivery order",
    "customs_declaration": "Customs declaration / release",
    "forwarder_invoice": "Forwarder invoice / charges",
    "delivery_proof": "Factory delivery proof",
}
EXPORT_DOCUMENTS = {
    "order": "Customer order / shipping instructions",
    "commercial_inv_pl": "Commercial invoice and packing list",
    "booking": "Carrier booking confirmation",
    "export_permit": "Export permit / certificates (if applicable)",
    "bill_of_lading": "Bill of lading",
    "customs_declaration": "Export customs declaration",
    "forwarder_invoice": "Forwarder invoice / charges",
    "delivery_proof": "Handover / proof of shipment",
}
DOCUMENT_TYPES = {**IMPORT_DOCUMENTS, **EXPORT_DOCUMENTS, "other": "Other supporting document"}
STATUSES = (
    "Preparing documents", "Documents under review", "Booked", "In transit",
    "At port", "Customs clearance", "Inland delivery", "Completed", "On hold",
)


@contextmanager
def _connect(database: Path):
    database.parent.mkdir(parents=True, exist_ok=True)
    connection = sqlite3.connect(database)
    connection.row_factory = sqlite3.Row
    connection.execute("PRAGMA foreign_keys = ON")
    connection.execute("PRAGMA busy_timeout = 5000")
    try:
        with connection:
            yield connection
    finally:
        connection.close()


def initialize(database: Path) -> None:
    with _connect(database) as connection:
        connection.executescript("""
            CREATE TABLE IF NOT EXISTS companies (
                id INTEGER PRIMARY KEY,
                name TEXT NOT NULL UNIQUE
            );
            CREATE TABLE IF NOT EXISTS shipments (
                id INTEGER PRIMARY KEY,
                company_id INTEGER NOT NULL REFERENCES companies(id),
                direction TEXT NOT NULL CHECK (direction IN ('Import', 'Export')),
                reference TEXT NOT NULL,
                partner TEXT NOT NULL DEFAULT '',
                origin TEXT NOT NULL DEFAULT '',
                destination TEXT NOT NULL DEFAULT '',
                hbl TEXT NOT NULL DEFAULT '',
                container TEXT NOT NULL DEFAULT '',
                forwarder TEXT NOT NULL DEFAULT '',
                etd TEXT NOT NULL DEFAULT '',
                eta_port TEXT NOT NULL DEFAULT '',
                eta_factory TEXT NOT NULL DEFAULT '',
                status TEXT NOT NULL DEFAULT 'Preparing documents',
                notes TEXT NOT NULL DEFAULT '',
                created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
                updated_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
                UNIQUE(company_id, direction, reference)
            );
            CREATE TABLE IF NOT EXISTS shipment_documents (
                id INTEGER PRIMARY KEY,
                shipment_id INTEGER NOT NULL REFERENCES shipments(id) ON DELETE CASCADE,
                kind TEXT NOT NULL,
                filename TEXT NOT NULL,
                relative_path TEXT NOT NULL,
                checksum TEXT NOT NULL,
                added_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
                UNIQUE(shipment_id, kind, checksum)
            );
            CREATE TABLE IF NOT EXISTS shipment_events (
                id INTEGER PRIMARY KEY,
                shipment_id INTEGER NOT NULL REFERENCES shipments(id) ON DELETE CASCADE,
                event_date TEXT NOT NULL,
                description TEXT NOT NULL,
                source TEXT NOT NULL DEFAULT '',
                created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
            );
            CREATE INDEX IF NOT EXISTS shipment_company_idx ON shipments(company_id, direction);
            CREATE INDEX IF NOT EXISTS document_shipment_idx ON shipment_documents(shipment_id);
        """)


def create_shipment(database: Path, company: str, direction: str, reference: str, **fields: str) -> int:
    company, reference = company.strip(), reference.strip()
    if not company or not reference or direction not in ("Import", "Export"):
        raise ValueError("Company, direction, and shipment reference are required.")
    allowed = {"partner", "origin", "destination", "hbl", "container", "forwarder", "etd", "eta_port", "eta_factory", "status", "notes"}
    if set(fields) - allowed:
        raise ValueError("Unknown shipment field.")
    if "status" in fields and fields["status"] not in STATUSES:
        raise ValueError("Unknown shipment status.")
    initialize(database)
    with _connect(database) as connection:
        connection.execute("INSERT OR IGNORE INTO companies(name) VALUES (?)", (company,))
        company_id = connection.execute("SELECT id FROM companies WHERE name=?", (company,)).fetchone()["id"]
        columns = ["company_id", "direction", "reference", *fields]
        values = [company_id, direction, reference, *fields.values()]
        placeholders = ",".join("?" for _ in columns)
        try:
            result = connection.execute(
                f"INSERT INTO shipments ({','.join(columns)}) VALUES ({placeholders})", values
            )
        except sqlite3.IntegrityError as exc:
            raise ValueError("This company already has a shipment with that direction and reference.") from exc
        return result.lastrowid


def update_shipment(database: Path, shipment_id: int, **fields: str) -> None:
    allowed = {"partner", "origin", "destination", "hbl", "container", "forwarder", "etd", "eta_port", "eta_factory", "status", "notes"}
    if not fields or set(fields) - allowed:
        raise ValueError("Unknown shipment field.")
    if "status" in fields and fields["status"] not in STATUSES:
        raise ValueError("Unknown shipment status.")
    setters = ", ".join(f"{key}=?" for key in fields)
    with _connect(database) as connection:
        result = connection.execute(
            f"UPDATE shipments SET {setters}, updated_at=CURRENT_TIMESTAMP WHERE id=?",
            [*fields.values(), shipment_id],
        )
        if not result.rowcount:
            raise ValueError("Shipment not found.")


def list_shipments(database: Path) -> list[dict]:
    initialize(database)
    with _connect(database) as connection:
        return [dict(row) for row in connection.execute("""
            SELECT s.*, c.name AS company FROM shipments s
            JOIN companies c ON c.id=s.company_id
            ORDER BY s.updated_at DESC, s.id DESC
        """)]


def get_shipment(database: Path, shipment_id: int) -> dict | None:
    with _connect(database) as connection:
        row = connection.execute("""
            SELECT s.*, c.name AS company FROM shipments s
            JOIN companies c ON c.id=s.company_id WHERE s.id=?
        """, (shipment_id,)).fetchone()
        return dict(row) if row else None


def _project_relative(project_root: Path, file_path: Path) -> str:
    try:
        return str(file_path.resolve().relative_to(project_root.resolve()))
    except ValueError as exc:
        raise ValueError("Document must be inside the project folder.") from exc


def attach_existing_document(database: Path, project_root: Path, shipment_id: int,
                             kind: str, file_path: Path) -> None:
    if kind not in DOCUMENT_TYPES:
        raise ValueError("Unknown document type.")
    if not file_path.is_file():
        raise FileNotFoundError(file_path)
    relative = _project_relative(project_root, file_path)
    checksum = sha256(file_path.read_bytes()).hexdigest()
    with _connect(database) as connection:
        connection.execute("""
            INSERT OR IGNORE INTO shipment_documents
            (shipment_id, kind, filename, relative_path, checksum) VALUES (?, ?, ?, ?, ?)
        """, (shipment_id, kind, file_path.name, relative, checksum))


def save_uploaded_document(database: Path, project_root: Path, shipment_id: int,
                           kind: str, filename: str, contents: bytes) -> Path:
    if kind not in DOCUMENT_TYPES or not contents:
        raise ValueError("Choose a document type and a non-empty file.")
    safe_name = re.sub(r"[^\w.() -]", "_", Path(filename).name, flags=re.UNICODE).strip(" .")
    if not safe_name or Path(safe_name).suffix.lower() not in {".pdf", ".xls", ".xlsx", ".doc", ".docx", ".jpg", ".jpeg", ".png"}:
        raise ValueError("Unsupported document filename or type.")
    checksum = sha256(contents).hexdigest()
    destination = project_root / "data" / "shipping_uploads" / str(shipment_id) / f"{checksum[:12]}_{safe_name}"
    destination.parent.mkdir(parents=True, exist_ok=True)
    if destination.exists() and sha256(destination.read_bytes()).hexdigest() != checksum:
        raise ValueError("A different file already exists at this path.")
    destination.write_bytes(contents)
    attach_existing_document(database, project_root, shipment_id, kind, destination)
    return destination


def list_documents(database: Path, shipment_id: int) -> list[dict]:
    with _connect(database) as connection:
        return [dict(row) for row in connection.execute(
            "SELECT * FROM shipment_documents WHERE shipment_id=? ORDER BY added_at, id", (shipment_id,)
        )]


def add_event(database: Path, shipment_id: int, event_date: str, description: str, source: str = "") -> None:
    if not description.strip():
        raise ValueError("Event description is required.")
    try:
        date.fromisoformat(event_date)
    except ValueError as exc:
        raise ValueError("Event date must be YYYY-MM-DD.") from exc
    with _connect(database) as connection:
        connection.execute("""
            INSERT INTO shipment_events (shipment_id, event_date, description, source)
            VALUES (?, ?, ?, ?)
        """, (shipment_id, event_date, description.strip(), source.strip()))


def list_events(database: Path, shipment_id: int) -> list[dict]:
    with _connect(database) as connection:
        return [dict(row) for row in connection.execute(
            "SELECT * FROM shipment_events WHERE shipment_id=? ORDER BY event_date DESC, id DESC", (shipment_id,)
        )]


def seed_tk164(database: Path, project_root: Path) -> None:
    """Add one clearly labeled starter record from the copied import register."""
    from openpyxl import load_workbook

    register = project_root / "reference/shipping_work/01_import_register/IMPORT LIST - 2026 ( XCF ).xlsx"
    if not register.is_file():
        return
    initialize(database)
    with _connect(database) as connection:
        existing = connection.execute("""
            SELECT s.id FROM shipments s JOIN companies c ON c.id=s.company_id
            WHERE c.name=? AND s.direction='Import' AND s.reference=?
        """, ("XIN CAI FU (CAMBODIA) CO., LTD.", "26TKYM164")).fetchone()
    if existing:
        return
    workbook = load_workbook(register, read_only=True, data_only=True)
    try:
        row = [cell.value for cell in workbook["ALL MONTHS"][6]]
    finally:
        workbook.close()
    if str(row[1]).strip() != "26TKYM164":
        raise ValueError("TK164 starter record does not match the source register.")
    def iso(value):
        return value.date().isoformat() if isinstance(value, datetime) else ""
    shipment_id = create_shipment(
        database, "XIN CAI FU (CAMBODIA) CO., LTD.", "Import", str(row[1]),
        partner=str(row[2] or ""), origin=str(row[8] or ""), destination="Sihanoukville, Cambodia",
        hbl=str(row[4] or ""), container=str(row[5] or ""), forwarder=str(row[22] or ""),
        etd=iso(row[9]), eta_port=iso(row[10]), eta_factory=iso(row[12]),
        status="Documents under review",
        notes="Starter record from copied 2026 import register. ETA dates are planned; actual arrival was blank in the source.",
    )
    base = project_root / "reference/shipping_work"
    for kind, relative in [
        ("supplier_inv_pl", "02_import_TK164/supplier_originals/清关---TK164拓卡-新彩福-柬埔寨资料.xls"),
        ("permit_inv_pl", "02_import_TK164/permit_application/1-26TKYM164.xlsx"),
        ("bill_of_lading", "02_import_TK164/transport/HBL电放-拆单_YS26100022-1.pdf"),
        ("arrival_notice", "02_import_TK164/transport/SITPSHSVA11652-SKH26IM10062.pdf"),
        ("forwarder_invoice", "02_import_TK164/forwarder_charges/TI-026-1231.pdf"),
    ]:
        file_path = base / relative
        if file_path.is_file():
            attach_existing_document(database, project_root, shipment_id, kind, file_path)
