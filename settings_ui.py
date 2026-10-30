"""Settings & Preferences Workspace: Default templates, CDC Master List quotas, and shipment register."""

from datetime import date
from pathlib import Path
import io
import pandas as pd
import streamlit as st
import openpyxl

from services.settings_store import (
    get_settings, save_settings,
    get_default_template_path, get_default_template_bytes, save_default_template,
    get_default_cdc_path, get_default_cdc_bytes, save_default_cdc,
)
from services.extraction import extract_cdc
from services.shipping_store import (
    DOCUMENT_TYPES, EXPORT_DOCUMENTS, IMPORT_DOCUMENTS, STATUSES, add_event,
    create_shipment, get_shipment, list_documents, list_events, list_shipments,
    save_uploaded_document, seed_tk164, update_shipment,
)


def render_settings_workspace(project_root: Path, lang: str = "en") -> None:
    """Render the Settings & Master Data preferences dashboard."""
    settings = get_settings()

    st.title("⚙️ Settings & Master Data Preferences")
    st.caption("Manage permanent default templates, CDC Master List quotas, company credentials, and track historical shipment registers.")

    tab_cdc, tab_tpl, tab_company, tab_shipments = st.tabs([
        "📊 CDC Master List & Quotas",
        "📄 Default Templates (INV / PKL)",
        "🏢 Company & Shipping Defaults",
        "📦 Shipment Register"
    ])

    # =========================================================================
    # TAB 1: CDC MASTER LIST & QUOTA OVERVIEW
    # =========================================================================
    with tab_cdc:
        st.markdown("### 📊 CDC Master List Setup & Overview")
        st.caption("Configure the active default CDC Master List file and review approved customs duty exemption quotas.")

        cdc_path = get_default_cdc_path()
        cdc_bytes = get_default_cdc_bytes()
        cdcmeta, cdcitems = extract_cdc(cdc_bytes, cdc_path.name)

        # 1. Active File Metrics Card
        with st.container(border=True):
            m_col1, m_col2, m_col3, m_col4 = st.columns(4)
            m_col1.metric("Registered Items", f"{len(cdcitems)} Items")
            co_name = cdcmeta.get("company") or settings.get("company_name", "—")
            m_col2.metric("Company Name", co_name[:25] + ("..." if len(co_name) > 25 else ""))
            qip_val = cdcmeta.get("qip") or settings.get("qip_reference", "—")
            m_col3.metric("QIP Reference", qip_val)
            file_kb = len(cdc_bytes) / 1024
            m_col4.metric("Active File", f"{cdc_path.name[:18]}... ({file_kb:.1f} KB)")

            b_c1, b_c2 = st.columns([1, 2], vertical_alignment="center")
            with b_c1:
                st.download_button(
                    "📥 Download Active Master List (.xlsx)",
                    data=cdc_bytes,
                    file_name=cdc_path.name,
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    width="stretch"
                )
            with b_c2:
                st.caption(f"📍 Stored at: `{cdc_path.relative_to(project_root)}`")

        # 2. Uploader to replace Default Master List
        with st.expander("📤 Upload & Set New Default CDC Master List (.xlsx)", expanded=False):
            st.caption("Upload a new official CDC Master List (.xlsx) to permanently replace the default master quota database.")
            uploaded_cdc = st.file_uploader(
                "Upload CDC Master List (.xlsx)",
                type=["xlsx"],
                key="settings_upload_cdc",
                label_visibility="collapsed"
            )
            if uploaded_cdc:
                new_bytes = uploaded_cdc.getvalue()
                try:
                    test_meta, test_items = extract_cdc(new_bytes, uploaded_cdc.name)
                    if not test_items:
                        st.error("No valid CDC quota items found in uploaded file. Please check worksheet structure.")
                    else:
                        st.success(f"✓ Valid CDC Master List detected! Found {len(test_items)} items for '{test_meta.get('company') or 'Registered Company'}'.")
                        if st.button("💾 Save as Permanent Default CDC Master List", type="primary"):
                            save_default_cdc(new_bytes, uploaded_cdc.name)
                            st.toast("CDC Master List updated successfully!", icon="✅")
                            st.rerun()
                except Exception as exc:
                    st.error(f"Error inspecting file: {exc}")

        st.markdown("<div style='height: 12px;'></div>", unsafe_allow_html=True)

        # 3. Overview Table of CDC Items
        st.markdown("#### 📋 Complete CDC Master Quota List")
        if not cdcitems:
            st.warning("No CDC items loaded in current master list.")
        else:
            filter_c1, filter_c2 = st.columns([2, 1], vertical_alignment="center")
            with filter_c1:
                search_query = st.text_input(
                    "Search items",
                    placeholder="🔍 Search by item name, CDC #, or unit...",
                    label_visibility="collapsed",
                    key="cdc_search_query"
                ).strip().lower()
            with filter_c2:
                cat_choice = st.pills(
                    "Category Filter",
                    options=["All", "Machinery (I)", "Raw Materials (II)"],
                    default="All",
                    label_visibility="collapsed",
                    key="cdc_cat_choice"
                )

            # Classify items into Category I (Machinery) vs II (Raw materials)
            table_rows = []
            for item in cdcitems:
                row_idx = item.get("row", 0)
                cat = "I - Machinery & Equipment" if row_idx < 100 else "II - Raw Materials & Accessories"
                s_str = str(item.get("serial", item.get("row", ""))).lstrip("#").strip()
                desc = item.get("description", "")
                unit = item.get("unit", "")
                appr = item.get("approved")
                price = item.get("unit_price")
                rem = item.get("remaining")

                # Filter logic
                if cat_choice == "Machinery (I)" and "Machinery" not in cat:
                    continue
                if cat_choice == "Raw Materials (II)" and "Raw Materials" not in cat:
                    continue
                if search_query:
                    combined = f"{s_str} {desc} {unit} {cat}".lower()
                    if search_query not in combined:
                        continue

                table_rows.append({
                    "CDC #": int(s_str) if s_str.isdigit() else s_str,
                    "Category": cat,
                    "Description": desc.upper(),
                    "Unit": unit,
                    "Approved Quota": f"{int(appr):,}" if appr is not None and float(appr).is_integer() else f"{appr:,.2f}" if appr is not None else "—",
                    "Unit Price (USD)": f"${price:,.2f}" if price is not None else "—",
                    "Remaining Balance": f"{int(rem):,}" if rem is not None and float(rem).is_integer() else f"{rem:,.2f}" if rem is not None else "—",
                })

            df_cdc = pd.DataFrame(table_rows)
            st.dataframe(df_cdc, hide_index=True, width="stretch")
            st.caption(f"Showing **{len(table_rows)}** of **{len(cdcitems)}** registered items in the official CDC Master List.")

    # =========================================================================
    # TAB 2: DEFAULT PERMIT TEMPLATES (INV / PKL)
    # =========================================================================
    with tab_tpl:
        st.markdown("### 📄 Default Permit Application Template (INV / PL)")
        st.caption("Inspect and manage the official 2-sheet customs permit template used to generate compliant INV and PL Excel files.")

        tpl_path = get_default_template_path()
        tpl_bytes = get_default_template_bytes()

        # 1. Template Overview Card
        with st.container(border=True):
            t_col1, t_col2, t_col3 = st.columns(3)
            t_col1.metric("Template Name", tpl_path.name)
            t_col2.metric("Worksheets", "INV (Invoice) + PL (Packing List)")
            t_col3.metric("Font & Format", "Times New Roman (Official)")

            b_t1, b_t2 = st.columns([1, 2], vertical_alignment="center")
            with b_t1:
                st.download_button(
                    "📥 Download Default Template (.xlsx)",
                    data=tpl_bytes,
                    file_name=tpl_path.name,
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    width="stretch"
                )
            with b_t2:
                st.caption(f"📍 Active path: `{tpl_path.relative_to(project_root)}`")

        # 2. Upload New Template
        with st.expander("📤 Upload & Set New Default Template (.xlsx)", expanded=False):
            st.caption("Upload a customized Permit Application workbook (.xlsx). Must contain an **INV** sheet and a **PL** sheet.")
            uploaded_tpl = st.file_uploader(
                "Upload Permit Template (.xlsx)",
                type=["xlsx"],
                key="settings_upload_tpl",
                label_visibility="collapsed"
            )
            if uploaded_tpl:
                tpl_new_bytes = uploaded_tpl.getvalue()
                try:
                    wb = openpyxl.load_workbook(io.BytesIO(tpl_new_bytes))
                    sheets_lower = [s.strip().lower() for s in wb.sheetnames]
                    has_inv = any("inv" in s for s in sheets_lower)
                    has_pl = any("pl" in s or "pack" in s for s in sheets_lower)

                    if not has_inv or not has_pl:
                        st.error(f"Workbook missing required sheets. Found: {wb.sheetnames}. Must include both INV and PL.")
                    else:
                        st.success(f"✓ Valid 2-sheet permit template detected! Sheets found: `{wb.sheetnames}`.")
                        if st.button("💾 Set as Permanent Default Template", type="primary"):
                            save_default_template(tpl_new_bytes, uploaded_tpl.name)
                            st.toast("Default permit template updated successfully!", icon="✅")
                            st.rerun()
                except Exception as exc:
                    st.error(f"Error inspecting template: {exc}")

        st.markdown("<div style='height: 12px;'></div>", unsafe_allow_html=True)

        # 3. Layout Structure Guide
        st.markdown("#### 📐 Standard Permit Layout Reference")
        with st.container(border=True):
            l_col1, l_col2 = st.columns(2)
            with l_col1:
                st.markdown("**Invoice Sheet (`INV`) Specifications:**")
                st.markdown("""
                - **Row 1-3:** Shipper name & address
                - **Row 8:** Invoice Number (Col B) & Issue Date (Col G)
                - **Row 9:** Shipping Mode (`BY TRUCK` / `BY SEA`)
                - **Row 10:** Route (`FROM ... TO ...`)
                - **Row 11:** Country of Origin (`CHINA`)
                - **Row 12-14:** Consignee & 2-Line Formatted Address
                - **Row 17:** VAT Number
                - **Row 23:** QIP No. & Category Header
                - **Row 24:** Items start (CDC #, Description, Qty, Unit, Unit Price, `=D{r}*F{r}`)
                - **Row 40:** Total Formula (`=SUM(G24:G33)`)
                """)
            with l_col2:
                st.markdown("**Packing List Sheet (`PL`) Specifications:**")
                st.markdown("""
                - **Row 1-3:** Shipper header (matches INV)
                - **Row 8:** Invoice No & Date evaluated static values (No raw `=INV!B8`)
                - **Row 9-11:** Shipping mode, Route, Country of origin
                - **Row 12-14:** Consignee & 2-Line Formatted Address
                - **Row 17:** VAT Number
                - **Row 22:** QIP No. & Category Header
                - **Row 23:** Items start (Packages, Description, Qty, Unit, Net Kg, Gross Kg)
                - **Row 39:** Total Packages Formula (`=SUM(A23:A32)`)
                - **Row 42:** Container Number & Port Info
                """)

    # =========================================================================
    # TAB 3: COMPANY & SHIPPING DEFAULTS
    # =========================================================================
    with tab_company:
        st.markdown("### 🏢 Company & Shipping Preferences")
        st.caption("Set default values used to auto-populate permit applications when not specified in shipping documents.")

        with st.form("company_defaults_form"):
            st.markdown("##### 🏛️ Company Registration & Tax")
            cf_col1, cf_col2, cf_col3 = st.columns(3)
            company_name = cf_col1.text_input("Registered Company Name", value=settings.get("company_name", ""))
            qip_ref = cf_col2.text_input("QIP File Reference", value=settings.get("qip_reference", "211/26"))
            vat_no = cf_col3.text_input("Tax / VAT Number", value=settings.get("vat_number", "LO01-902303477"))

            st.markdown("##### 📍 Consignee Address (2-Line Permit Format)")
            addr_col1, addr_col2 = st.columns(2)
            addr_1 = addr_col1.text_input("Address Line 1 (Sangkat / Commune)", value=settings.get("consignee_address_1", "PHUM DOMNAK TROYEOUNG, SANGKAT CHOM CHAO II,"))
            addr_2 = addr_col2.text_input("Address Line 2 (Khan / City)", value=settings.get("consignee_address_2", "KHAN POR SEN CHEY PHNOM PENH,"))
            addr_3 = st.text_input("Country Line", value=settings.get("consignee_address_3", "KINGDOM OF CAMBODIA."))

            st.markdown("##### 🚢 Default Shipping & Route")
            sr_col1, sr_col2, sr_col3 = st.columns(3)
            modes = ["BY TRUCK", "BY SEA", "BY AIR"]
            cur_mode = settings.get("shipping_mode", "BY TRUCK")
            mode_idx = modes.index(cur_mode) if cur_mode in modes else 0
            ship_mode = sr_col1.selectbox("Default Shipping Mode", modes, index=mode_idx)
            origin = sr_col2.text_input("Default Origin Port", value=settings.get("origin", "JIANGSU, CHINA"))
            dest = sr_col3.text_input("Default Destination Port", value=settings.get("destination", "PHNOM PENH, CAMBODIA"))

            saved_pref = st.form_submit_button("💾 Save Company & Shipping Preferences", type="primary")

        if saved_pref:
            updated = {
                "company_name": company_name.strip(),
                "qip_reference": qip_ref.strip(),
                "vat_number": vat_no.strip(),
                "consignee_address_1": addr_1.strip(),
                "consignee_address_2": addr_2.strip(),
                "consignee_address_3": addr_3.strip(),
                "shipping_mode": ship_mode,
                "origin": origin.strip(),
                "destination": dest.strip(),
                "template_filename": settings.get("template_filename", "permit_template_default.xlsx"),
                "cdc_filename": settings.get("cdc_filename", "default_cdc_master.xlsx"),
            }
            save_settings(updated)
            st.toast("Company & Shipping Preferences saved!", icon="💾")
            st.success("✓ Preferences saved successfully. All future permit generations will adopt these defaults.")

    # =========================================================================
    # TAB 4: SHIPMENT REGISTER (HISTORICAL TRACKER)
    # =========================================================================
    with tab_shipments:
        st.markdown("### 📦 Shipment Register")
        st.caption("Internal record tracking for documents, forwarder transport, and milestones across imports and exports.")

        database = project_root / "data" / "shipping.sqlite3"
        seed_tk164(database, project_root)

        with st.expander("➕ Create New Shipment Record", expanded=False):
            with st.form("create_shipment_form", clear_on_submit=True):
                c1, c2, c3 = st.columns(3)
                c_name = c1.text_input("Company", value=settings.get("company_name", "XIN CAI FU (CAMBODIA) CO., LTD."))
                direction = c2.selectbox("Direction", ["Import", "Export"])
                reference = c3.text_input("Shipment Reference", placeholder="e.g. 26TKYM164")
                partner = st.text_input("Supplier / Partner Name", placeholder="e.g. SHANGHAI TOKAYMA...")
                submitted = st.form_submit_button("Create Shipment", type="primary")
            if submitted:
                try:
                    shipment_id = create_shipment(database, c_name, direction, reference, partner=partner)
                    st.session_state["selected_shipment_id"] = shipment_id
                    st.toast("Shipment record created!", icon="📦")
                    st.rerun()
                except ValueError as exc:
                    st.error(str(exc))

        shipments = list_shipments(database)
        if not shipments:
            st.info("No shipments created yet.")
            return

        filter_col1, filter_col2 = st.columns(2)
        companies = sorted({item["company"] for item in shipments})
        company_filter = filter_col1.selectbox("Filter by Company", ["All companies", *companies])
        direction_filter = filter_col2.selectbox("Filter by Direction", ["All directions", "Import", "Export"])
        visible = [
            item for item in shipments
            if (company_filter == "All companies" or item["company"] == company_filter)
            and (direction_filter == "All directions" or item["direction"] == direction_filter)
        ]

        st.dataframe(pd.DataFrame([{
            "Company": item["company"], "Direction": item["direction"],
            "Reference": item["reference"], "Status": item["status"],
            "HBL": item["hbl"], "Container": item["container"],
            "ETA port": item["eta_port"], "ETA factory": item["eta_factory"],
        } for item in visible]), hide_index=True, width="stretch")

        if not visible:
            st.info("No shipments match these filters.")
            return

        ids = [item["id"] for item in visible]
        saved = st.session_state.get("selected_shipment_id")
        selection = st.selectbox(
            "Select Shipment to Manage", ids, index=ids.index(saved) if saved in ids else 0,
            format_func=lambda shipment_id: next(
                f"{item['company']} · {item['direction']} · {item['reference']}"
                for item in visible if item["id"] == shipment_id
            ),
        )
        st.session_state["selected_shipment_id"] = selection
        shipment = get_shipment(database, selection)

        st.divider()
        st.subheader(f"{shipment['reference']} · {shipment['direction']}")
        st.write(f"**Company:** {shipment['company']}  ·  **Partner:** {shipment['partner'] or '—'}")

        with st.form(f"details_{selection}"):
            c1, c2, c3 = st.columns(3)
            status = c1.selectbox("Status", STATUSES, index=STATUSES.index(shipment["status"]))
            hbl = c2.text_input("B/L or HBL Number", value=shipment["hbl"])
            container = c3.text_input("Container Number", value=shipment["container"])
            c4, c5, c6 = st.columns(3)
            orig_val = c4.text_input("Origin", value=shipment["origin"])
            dest_val = c5.text_input("Destination", value=shipment["destination"])
            fwd_val = c6.text_input("Forwarder", value=shipment["forwarder"])
            c7, c8, c9 = st.columns(3)
            etd = c7.text_input("ETD (YYYY-MM-DD)", value=shipment["etd"])
            eta_port = c8.text_input("ETA Port (YYYY-MM-DD)", value=shipment["eta_port"])
            eta_factory = c9.text_input("ETA Factory (YYYY-MM-DD)", value=shipment["eta_factory"])
            notes = st.text_area("Notes", value=shipment["notes"])
            save = st.form_submit_button("Save Shipment Details")

        if save:
            try:
                for value in (etd, eta_port, eta_factory):
                    if value:
                        date.fromisoformat(value)
                update_shipment(
                    database, selection, status=status, hbl=hbl.strip(),
                    container=container.strip(), origin=orig_val.strip(),
                    destination=dest_val.strip(), forwarder=fwd_val.strip(),
                    etd=etd, eta_port=eta_port, eta_factory=eta_factory, notes=notes.strip()
                )
                st.toast("Shipment updated successfully!", icon="✅")
            except ValueError as exc:
                st.error(f"Check the dates and fields: {exc}")

        st.markdown("##### 📁 Attached Documents")
        documents = list_documents(database, selection)
        checklist = IMPORT_DOCUMENTS if shipment["direction"] == "Import" else EXPORT_DOCUMENTS
        rows = []
        for kind, label in checklist.items():
            attached = [item for item in documents if item["kind"] == kind]
            rows.append({
                "Document": label,
                "Status": "Attached" if attached else "Missing / Not Recorded",
                "Files": ", ".join(item["filename"] for item in attached)
            })
        st.dataframe(pd.DataFrame(rows), hide_index=True, width="stretch")

        with st.form(f"attach_{selection}"):
            kind = st.selectbox(
                "Document Type", list(checklist) + ["other"],
                format_func=lambda value: DOCUMENT_TYPES.get(value, value)
            )
            upload = st.file_uploader("Attach Document", type=["pdf", "xls", "xlsx", "doc", "docx", "jpg", "jpeg", "png"])
            attach = st.form_submit_button("Save Attachment")

        if attach:
            if upload is None:
                st.error("Choose a file first.")
            else:
                try:
                    save_uploaded_document(database, project_root, selection, kind, upload.name, upload.getvalue())
                    st.toast("Document attached successfully!", icon="📎")
                    st.rerun()
                except (OSError, ValueError) as exc:
                    st.error(str(exc))
